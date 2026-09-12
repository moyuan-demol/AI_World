"""Document loaders: PDF / DOCX / TXT / MD -> plain text."""

import io
import logging
from pathlib import Path

from app.core.errors import ValidationError

logger = logging.getLogger(__name__)

TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".text",
    ".csv",
    ".json",
    ".html",
    ".htm",
}
XLSX_EXTENSIONS = {".xlsx", ".xlsm"}
PPTX_EXTENSIONS = {".pptx"}
PDF_EXTENSIONS = {".pdf"}
DOCX_EXTENSIONS = {".docx"}
SUPPORTED_EXTENSIONS = (
    TEXT_EXTENSIONS | PDF_EXTENSIONS | DOCX_EXTENSIONS | XLSX_EXTENSIONS | PPTX_EXTENSIONS
)

_ENCODINGS = ("utf-8", "utf-8-sig", "gb18030", "big5", "latin-1")


def load_from_path(path: str | Path) -> str:
    path = Path(path)
    return load_from_bytes(path.read_bytes(), path.name)


def load_from_bytes(data: bytes, filename: str) -> str:
    suffix = Path(filename or "").suffix.lower()
    if suffix in TEXT_EXTENSIONS:
        return _decode(data)
    if suffix in PDF_EXTENSIONS:
        return _load_pdf(data)
    if suffix in DOCX_EXTENSIONS:
        return _load_docx(data)
    if suffix in XLSX_EXTENSIONS:
        return _load_xlsx(data)
    if suffix in PPTX_EXTENSIONS:
        return _load_pptx(data)
    raise ValidationError(
        "暂不支持 " + (suffix or "该") + " 格式，支持：" + ", ".join(sorted(SUPPORTED_EXTENSIONS))
    )


def _decode(data: bytes) -> str:
    for encoding in _ENCODINGS:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _load_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise ValidationError("解析 PDF 需要安装 pypdf，请执行 pip install -r requirements.txt") from exc

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise ValidationError("PDF 文件已损坏或无法读取：" + str(exc)) from exc

    pages: list[str] = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception as exc:
            logger.warning("PDF 页面解析失败: %s", exc)
    text = "\n".join(pages).strip()
    if not text:
        raise ValidationError("未能从 PDF 中提取到文本（可能是扫描件，需要 OCR）")
    return text


def _load_xlsx(data: bytes) -> str:
    """Excel：每个工作表按「单元格 | 单元格」逐行导出为文本。"""
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise ValidationError("解析 XLSX 需要安装 openpyxl") from exc
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:
        raise ValidationError("XLSX 文件已损坏或无法读取：" + str(exc)) from exc
    parts: list[str] = []
    for sheet in workbook.worksheets:
        parts.append("# 工作表：" + str(sheet.title))
        for row in sheet.iter_rows(values_only=True):
            cells = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
            if cells:
                parts.append(" | ".join(cells))
    workbook.close()
    text = "\n".join(parts).strip()
    if not text:
        raise ValidationError("未能从 XLSX 中提取到文本")
    return text


def _load_pptx(data: bytes) -> str:
    """PPT：导出每页的标题与正文。"""
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover
        raise ValidationError("解析 PPTX 需要安装 python-pptx") from exc
    try:
        presentation = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise ValidationError("PPTX 文件已损坏或无法读取：" + str(exc)) from exc
    parts: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        parts.append("# 第 " + str(index) + " 页")
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if text and text.strip():
                parts.append(text.strip())
    text = "\n".join(parts).strip()
    if not text:
        raise ValidationError("未能从 PPTX 中提取到文本")
    return text


def _load_docx(data: bytes) -> str:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise ValidationError("解析 DOCX 需要安装 python-docx，请执行 pip install -r requirements.txt") from exc

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ValidationError("DOCX 文件已损坏或无法读取：" + str(exc)) from exc

    parts: list[str] = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts).strip()
    if not text:
        raise ValidationError("未能从 DOCX 中提取到文本")
    return text
